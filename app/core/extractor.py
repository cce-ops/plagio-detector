import io
from pathlib import Path

import fitz  # pymupdf
from docx import Document
from PIL import Image

# OCR: solo si la pagina no tiene texto extraible, y con limites de coste
OCR_DPI = 200
OCR_MAX_PAGINAS = 40          # evita análisis de 100 paginas escaneadas
OCR_TIMEOUT_GLOBAL = 180      # segundos maxima de OCR por documento

TESSERACT_CMD = r"C:\Program Files\Tesseract-OCR\tesseract.exe"


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix in (".docx", ".doc"):
        return _extract_docx(path)
    if suffix in (".txt", ".md"):
        return path.read_text(encoding="utf-8", errors="ignore")
    if suffix in (".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff"):
        return _extract_image(path)
    raise ValueError(f"Formato no soportado: {suffix}")


def _extract_pdf(path: Path) -> str:
    doc = fitz.open(path)
    textos: list[str] = []
    import time

    inicio = time.time()

    for i, page in enumerate(doc):
        texto = page.get_text().strip()
        if texto:
            textos.append(texto)
            continue

        # Pagina sin texto: probablemente escaneada. Solo si no superamos limites.
        if i >= OCR_MAX_PAGINAS:
            textos.append(f"[pagina {i + 1}: sin texto extraible, OCR omitido por limite]")
            continue
        if time.time() - inicio > OCR_TIMEOUT_GLOBAL:
            textos.append(f"[pagina {i + 1}: sin texto extraible, OCR omitido por tiempo]")
            continue

        try:
            pix = page.get_pixmap(dpi=OCR_DPI)
            img = Image.open(io.BytesIO(pix.tobytes("png")))
            ocr = _ocr_image(img)
            textos.append(ocr if ocr.strip() else f"[pagina {i + 1}: sin texto]")
        except Exception as e:
            textos.append(f"[pagina {i + 1}: error OCR {type(e).__name__}]")

    doc.close()
    return "\n".join(textos)


def _extract_docx(path: Path) -> str:
    doc = Document(path)
    partes = [p.text for p in doc.paragraphs if p.text.strip()]

    # Tablas: contienen los datos numericos del proyecto
    for tabla in doc.tables:
        for fila in tabla.rows:
            celdas = [c.text.strip() for c in fila.cells if c.text.strip()]
            if celdas:
                partes.append(" | ".join(celdas))

    return "\n".join(partes)


def _extract_image(path: Path) -> str:
    return _ocr_image(Image.open(path))


def _ocr_image(img: Image.Image) -> str:
    try:
        import pytesseract
    except ImportError:
        return "[pytesseract no instalado: falta OCR]"
    try:
        pytesseract.pytesseract.tesseract_cmd = TESSERACT_CMD
        return pytesseract.image_to_string(img, lang="spa")
    except Exception as e:
        return f"[error OCR: {type(e).__name__}: {e}]"